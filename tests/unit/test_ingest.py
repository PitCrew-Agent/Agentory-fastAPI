"""매뉴얼 파싱·정규화 단위 테스트 (AI_RAG01_PREP01)

정규화는 인라인 문자열로 결정론적 검증, 파서는 tmp_path 사용
DOCX 파싱은 Docling 변환 경유, 문단은 마크다운 규칙(빈 줄)으로 구분
특수문자는 이스케이프 대신 chr로 구성해 소스 가독성·정확성 확보
"""

import unicodedata

import pytest
from docx import Document

from agentory.modules.rag.ingest import (
    SUPPORTED_SUFFIXES,
    build_chunks,
    chunk_document,
    normalize,
    parse,
    parse_and_normalize,
)

NL = chr(10)  # 개행
CR = chr(13)  # 캐리지리턴
TAB = chr(9)  # 탭
SOFT_HYPHEN = chr(0x00AD)  # 소프트하이픈
ZWSP = chr(0x200B)  # 제로폭 공백
NUL = chr(0)  # 널 문자
BEL = chr(7)  # 제어문자 예시


def test_normalize_unifies_crlf_and_cr_to_lf():
    assert normalize("a" + CR + NL + "b" + CR + "c") == "a" + NL + "b" + NL + "c"


def test_normalize_composes_hangul_to_nfc():
    decomposed = unicodedata.normalize("NFD", "설비")
    result = normalize(decomposed)
    assert unicodedata.is_normalized("NFC", result)
    assert result == "설비"


def test_normalize_removes_soft_hyphen_and_zero_width():
    assert normalize("tem" + SOFT_HYPHEN + "per" + ZWSP + "ature") == "temperature"


def test_normalize_strips_control_chars():
    assert normalize("a" + NUL + "b" + BEL + "c") == "abc"


def test_normalize_keeps_newline_and_collapses_tab_to_space():
    # 제어문자 제거 단계는 탭·개행 보존, 공백 축약 단계에서 탭은 단일 공백
    assert normalize("a" + TAB + "b" + NL + "c") == "a b" + NL + "c"


def test_normalize_dehyphenates_latin_word_across_linebreak():
    assert normalize("cool-" + NL + "down") == "cooldown"


def test_normalize_preserves_hyphen_in_codes_and_names():
    assert normalize("ERR-402 A-Line") == "ERR-402 A-Line"


def test_normalize_collapses_repeated_spaces_and_tabs():
    assert normalize("a    b" + TAB + "c") == "a b c"


def test_normalize_strips_trailing_whitespace_per_line():
    assert normalize("a  " + NL + "b  ") == "a" + NL + "b"


def test_normalize_collapses_excess_blank_lines():
    assert normalize("a" + NL * 4 + "b") == "a" + NL * 2 + "b"


def test_normalize_is_idempotent():
    messy = "제목" + CR + NL * 4 + "본문  " + TAB + "줄" + NUL + NL + "cool-" + NL + "down"
    once = normalize(messy)
    assert normalize(once) == once


def test_normalize_empty_returns_empty():
    assert normalize("") == ""


def test_normalize_whitespace_only_returns_empty():
    assert normalize("   " + NL + TAB + NL + "  ") == ""


def test_parse_reads_txt(tmp_path):
    path = tmp_path / "manual.txt"
    path.write_text("센서 점검 절차", encoding="utf-8")
    assert parse(path) == "센서 점검 절차"


def test_parse_reads_md(tmp_path):
    path = tmp_path / "manual.md"
    path.write_text("# 제목" + NL + "본문", encoding="utf-8")
    assert parse(path) == "# 제목" + NL + "본문"


def test_parse_reads_docx_via_docling(tmp_path):
    # Docling 변환은 문단을 마크다운 규칙(빈 줄)으로 구분
    path = tmp_path / "manual.docx"
    document = Document()
    document.add_paragraph("첫 문단")
    document.add_paragraph("둘째 문단")
    document.save(str(path))
    assert parse(path) == "첫 문단" + NL * 2 + "둘째 문단"


def test_parse_rejects_unsupported_suffix(tmp_path):
    path = tmp_path / "manual.doc"
    path.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        parse(path)


def test_parse_and_normalize_matches_normalize_of_parse(tmp_path):
    path = tmp_path / "manual.txt"
    path.write_text("제목  " + TAB + NL * 4 + "본문", encoding="utf-8")
    assert parse_and_normalize(path) == normalize(parse(path))


def test_supported_suffixes_are_expected_formats():
    assert SUPPORTED_SUFFIXES == {".pdf", ".docx", ".txt", ".md"}


# --- 청킹 (AI_RAG01_CHUNK01) ---


def test_chunk_document_md_via_docling_returns_nonempty_chunks(tmp_path):
    # md는 Docling 변환 후 HybridChunker로 헤딩·표 구조 인식 분할
    path = tmp_path / "manual.md"
    md = (
        "# 에칭 설비"
        + NL * 2
        + "## ERR-402 냉각 이상"
        + NL * 2
        + "냉각수 유량 점검 후 밸브 교체"
        + NL
    )
    path.write_text(md, encoding="utf-8")
    chunks = chunk_document(path)
    assert chunks
    assert all(isinstance(piece, str) and piece for piece in chunks)
    assert "냉각수 유량 점검 후 밸브 교체" in " ".join(chunks)


def test_chunk_document_txt_token_splits(tmp_path):
    # txt는 Docling 미지원이라 토큰 창 분할, 작은 max_tokens로 다중 청크 확인
    path = tmp_path / "manual.txt"
    path.write_text("가나다 " * 200, encoding="utf-8")
    chunks = chunk_document(path, max_tokens=16)
    assert len(chunks) > 1
    assert all(piece for piece in chunks)


def test_chunk_document_rejects_unsupported_suffix(tmp_path):
    path = tmp_path / "manual.doc"
    path.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        chunk_document(path)


def test_chunk_document_rejects_nonpositive_max_tokens(tmp_path):
    path = tmp_path / "manual.txt"
    path.write_text("본문", encoding="utf-8")
    with pytest.raises(ValueError):
        chunk_document(path, max_tokens=0)


def test_build_chunks_assigns_sequential_index_and_metadata():
    chunks = build_chunks(
        ["첫 청크", "둘째 청크", "셋째 청크"],
        doc_id="MAN-TST-001",
        equipment_type="Etching",
        alarm_code="ERR-402",
    )
    assert [c["chunk_index"] for c in chunks] == [0, 1, 2]
    assert all(c["doc_id"] == "MAN-TST-001" for c in chunks)
    assert all(c["equipment_type"] == "Etching" for c in chunks)
    assert all(c["alarm_code"] == "ERR-402" for c in chunks)
    assert all(c["content"] for c in chunks)
    assert all("embedding" not in c for c in chunks)


def test_build_chunks_defaults_metadata_to_none():
    chunks = build_chunks(["abc"], doc_id="MAN-TST-002")
    assert chunks[0]["equipment_type"] is None
    assert chunks[0]["alarm_code"] is None


def test_build_chunks_empty_list_returns_empty_list():
    assert build_chunks([], doc_id="MAN-TST-003") == []
